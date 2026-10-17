# -*- coding: utf-8 -*-
"""在多图版Excel基础上追加3列AI生成图片。"""
import json, os, urllib.request, hashlib
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AI_IMG = json.load(open(os.path.join(ROOT, "data/normalized/heritage_ai_images.json"), encoding="utf-8"))
CACHE_DIR = os.path.join(ROOT, "static/heritage_ai_cache")
SRC = os.path.join(ROOT, "data/normalized/山东省国家级非遗项目清单_多图版.xlsx")
OUT = os.path.join(ROOT, "data/normalized/山东省国家级非遗项目清单_最终版.xlsx")
os.makedirs(CACHE_DIR, exist_ok=True)

THUMB_W = 160
THUMB_H = 120
NUM_AI_COLS = 3
ROW_HEIGHT = 100

def download_img(url):
    h = hashlib.md5(url.encode()).hexdigest()[:16]
    p = os.path.join(CACHE_DIR, f"{h}.jpg")
    if os.path.exists(p) and os.path.getsize(p) > 1000:
        return p
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15) as r:
            data = r.read()
        with open(p, "wb") as f:
            f.write(data)
        return p
    except Exception as e:
        print(f"  download fail: {url[:60]} -> {e}")
        return None

wb = openpyxl.load_workbook(SRC)
ws = wb["项目明细"]

# 现有列: A-H(8列数据) + I-M(5列真实图片) = 13列
# 新增: N, O, P 三列AI配图
start_col = ws.max_column + 1  # 14

hdr_font = Font(bold=True, size=11, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="C00000")  # 红色表头区分AI图
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# 写3个AI配图表头
for i in range(NUM_AI_COLS):
    col_idx = start_col + i
    col_letter = get_column_letter(col_idx)
    c = ws.cell(row=1, column=col_idx, value=f"AI配图{i+1}")
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = center
    c.border = border
    ws.column_dimensions[col_letter].width = 20

# 遍历数据行，匹配项目名，嵌入AI图
embedded = 0
missing = 0
for row_num in range(2, ws.max_row + 1):
    ws.row_dimensions[row_num].height = ROW_HEIGHT
    name = ws.cell(row=row_num, column=2).value  # B列=项目名称
    if not name:
        continue

    # AI图URL列表（每个项目3张）
    ai_urls = AI_IMG.get(name, [])

    for img_idx in range(NUM_AI_COLS):
        col_idx = start_col + img_idx
        col_letter = get_column_letter(col_idx)
        ws.cell(row=row_num, column=col_idx).border = border
        ws.cell(row=row_num, column=col_idx).alignment = Alignment(vertical="center")

        if img_idx < len(ai_urls) and ai_urls[img_idx]:
            local = download_img(ai_urls[img_idx])
            if local and os.path.exists(local):
                try:
                    pil = PILImage.open(local)
                    pil.thumbnail((THUMB_W, THUMB_H))
                    tmp = local.replace(".jpg", f"_a{img_idx}.png")
                    pil.save(tmp, "PNG")
                    xi = XLImage(tmp)
                    xi.width = THUMB_W
                    xi.height = THUMB_H
                    ws.add_image(xi, f"{col_letter}{row_num}")
                    embedded += 1
                except Exception as e:
                    ws.cell(row=row_num, column=col_idx).value = "生成失败"

    if not ai_urls:
        missing += 1

wb.save(OUT)
print(f"Saved: {OUT}")
print(f"AI images embedded: {embedded}, rows missing AI: {missing}")

# 清理临时文件
for f in os.listdir(CACHE_DIR):
    if "_a" in f and f.endswith(".png"):
        try:
            os.remove(os.path.join(CACHE_DIR, f))
        except Exception:
            pass
