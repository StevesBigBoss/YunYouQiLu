import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV = os.path.join(ROOT, "data", "normalized", "heritage_table.csv")
OUT = os.path.join(ROOT, "data", "normalized", "山东省国家级非遗项目清单.xlsx")

rows = list(csv.DictReader(open(CSV, encoding="utf-8-sig")))

wb = Workbook()
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="305496")
head_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
body_font = Font(name="微软雅黑", size=10)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

ws = wb.active
ws.title = "项目明细"
cols = ["序号", "项目名称", "类别", "批次", "年份", "市", "县/区", "保护单位"]
ws.append(cols)
for c in range(1, len(cols)+1):
    cell = ws.cell(row=1, column=c)
    cell.fill = head_fill; cell.font = head_font; cell.alignment = center; cell.border = border

for i, r in enumerate(rows, 1):
    ws.append([i, r["项目名称"], r["类别"], r["批次"], int(r["年份"]) if r["年份"] else "",
               r["市"], r["县/区"], r["保护单位"]])

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(cols)):
    for cell in row:
        cell.font = body_font; cell.border = border
        cell.alignment = left if cell.column in (2, 8) else center

widths = [6, 32, 18, 10, 8, 18, 16, 42]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

ws2 = wb.create_sheet("按市汇总")
from collections import Counter
city_ct = Counter(r["市"] for r in rows)
ws2.append(["市", "项目数"])
for c in (1, 2):
    cell = ws2.cell(row=1, column=c)
    cell.fill = head_fill; cell.font = head_font; cell.alignment = center; cell.border = border
for city, n in sorted(city_ct.items(), key=lambda x: -x[1]):
    ws2.append([city, n])
ws2.append(["合计", len(rows)])
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=2):
    for cell in row:
        cell.font = body_font; cell.border = border; cell.alignment = center
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 10

ws3 = wb.create_sheet("按类别汇总")
cat_ct = Counter(r["类别"] for r in rows)
ws3.append(["类别", "项目数"])
for c in (1, 2):
    cell = ws3.cell(row=1, column=c)
    cell.fill = head_fill; cell.font = head_font; cell.alignment = center; cell.border = border
for cat, n in sorted(cat_ct.items(), key=lambda x: -x[1]):
    ws3.append([cat, n])
ws3.append(["合计", len(rows)])
for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, max_col=2):
    for cell in row:
        cell.font = body_font; cell.border = border; cell.alignment = center
ws3.column_dimensions["A"].width = 24
ws3.column_dimensions["B"].width = 10

wb.save(OUT)
print("saved:", OUT, "rows:", len(rows))
