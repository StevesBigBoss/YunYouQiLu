# -*- coding: utf-8 -*-
"""生成带5列多图的最终 Excel。"""
import json, os, urllib.request, hashlib
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERITAGE = json.load(open(os.path.join(ROOT, "data/normalized/heritage.json"), encoding="utf-8"))
MULTI_IMG = json.load(open(os.path.join(ROOT, "data/normalized/heritage_multi_images.json"), encoding="utf-8"))
SINGLE_IMG = json.load(open(os.path.join(ROOT, "data/normalized/heritage_image_map.json"), encoding="utf-8"))
CACHE_DIR = os.path.join(ROOT, "static/heritage_cache")
OUT = os.path.join(ROOT, "data/normalized/山东省国家级非遗项目清单_多图版.xlsx")
os.makedirs(CACHE_DIR, exist_ok=True)

# 从现有 Excel 读市/县数据
OLD_OUT = os.path.join(ROOT, "data/normalized/山东省国家级非遗项目清单.xlsx")
old_wb = openpyxl.load_workbook(OLD_OUT)
old_ws = old_wb["项目明细"]
rows_data = []
for row in old_ws.iter_rows(min_row=2, values_only=True):
    rows_data.append({
        "seq": row[0], "name": row[1], "category": row[2],
        "batch": row[3], "year": row[4], "city": row[5],
        "district": row[6], "unit": row[7]
    })

# 建立 items 映射
items_by_name = {}
for it in HERITAGE["items"]:
    items_by_name.setdefault(it["name"], []).append(it)

THUMB_W = 90
THUMB_H = 68
ROW_HEIGHT = 55
NUM_IMG_COLS = 5

def url_to_cache_path(url):
    h = hashlib.md5(url.encode()).hexdigest()[:16]
    return os.path.join(CACHE_DIR, f"{h}.jpg")

def download_img(url):
    p = url_to_cache_path(url)
    if os.path.exists(p) and os.path.getsize(p) > 1000:
        return p
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15) as r:
            data = r.read()
        with open(p, "wb") as f:
            f.write(data)
        return p
    except Exception as e:
        print(f"  download fail: {url[:60]} -> {e}")
        return None

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "项目明细"

# 表头: 序号, 项目名称, 类别, 批次, 年份, 市, 县/区, 保护单位, 图片1..图片5
headers = ["序号", "项目名称", "类别", "批次", "年份", "市", "县/区", "保护单位"]
for i in range(1, NUM_IMG_COLS + 1):
    headers.append(f"图片{i}")
ws.append(headers)

# 样式
hdr_font = Font(bold=True, size=11, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="305496")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col_idx in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col_idx)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = center
    c.border = border

# 列宽
widths = [6, 30, 12, 10, 8, 10, 14, 36] + [12] * NUM_IMG_COLS
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

total_imgs = 0
for idx, rd in enumerate(rows_data):
    row_num = idx + 2
    row_vals = [rd["seq"], rd["name"], rd["category"], rd["batch"],
                rd["year"], rd["city"], rd["district"], rd["unit"]]
    # 图片列先填空字符串
    row_vals += [""] * NUM_IMG_COLS
    ws.append(row_vals)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=row_num, column=col)
        c.border = border
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row_num].height = ROW_HEIGHT

    name = rd["name"]
    # 获取5张图URL
    img_urls = MULTI_IMG.get(name, [])
    if not img_urls:
        single = SINGLE_IMG.get(name)
        if single:
            img_urls = [single]

    # 下载并嵌入
    for img_idx, url in enumerate(img_urls[:NUM_IMG_COLS]):
        col_letter = get_column_letter(8 + img_idx + 1)  # 第9列起
        local = download_img(url)
        if local and os.path.exists(local):
            try:
                pil = PILImage.open(local)
                pil.thumbnail((THUMB_W, THUMB_H))
                tmp_path = local.replace(".jpg", f"_t{img_idx}.png")
                pil.save(tmp_path, "PNG")
                xi = XLImage(tmp_path)
                xi.width = THUMB_W
                xi.height = THUMB_H
                ws.add_image(xi, f"{col_letter}{row_num}")
                total_imgs += 1
            except Exception as e:
                ws.cell(row=row_num, column=8 + img_idx + 1).value = "图片错误"

ws.freeze_panes = "A2"

# Sheet2 按市汇总
ws2 = wb.create_sheet("按市汇总")
city_count = {}
for rd in rows_data:
    city_count[rd["city"]] = city_count.get(rd["city"], 0) + 1
ws2.append(["市", "项目数"])
for c, n in sorted(city_count.items(), key=lambda x: -x[1]):
    ws2.append([c, n])
for col in [1, 2]:
    ws2.cell(row=1, column=col).font = hdr_font
    ws2.cell(row=1, column=col).fill = hdr_fill
ws2.column_dimensions["A"].width = 14
ws2.column_dimensions["B"].width = 10

# Sheet3 按类别汇总
ws3 = wb.create_sheet("按类别汇总")
cat_count = {}
for rd in rows_data:
    cat_count[rd["category"]] = cat_count.get(rd["category"], 0) + 1
ws3.append(["类别", "项目数"])
for c, n in sorted(cat_count.items(), key=lambda x: -x[1]):
    ws3.append([c, n])
for col in [1, 2]:
    ws3.cell(row=1, column=col).font = hdr_font
    ws3.cell(row=1, column=col).fill = hdr_fill
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 10

wb.save(OUT)
print(f"Saved: {OUT}")
print(f"Rows: {len(rows_data)}, total images embedded: {total_imgs}")

# 清理临时缩略图
for f in os.listdir(CACHE_DIR):
    if "_t" in f and f.endswith(".png"):
        try:
            os.remove(os.path.join(CACHE_DIR, f))
        except Exception:
            pass
