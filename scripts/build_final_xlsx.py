import json, os
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERITAGE = json.load(open(os.path.join(ROOT, "data/normalized/heritage.json"), encoding="utf-8"))
IMG_DIR = os.path.join(ROOT, "static/heritage")
OUT = os.path.join(ROOT, "data/normalized/山东省国家级非遗项目清单.xlsx")

old_wb = openpyxl.load_workbook(OUT)
old_ws = old_wb["项目明细"]
rows_data = []
for row in old_ws.iter_rows(min_row=2, values_only=True):
    rows_data.append({
        "seq": row[0], "name": row[1], "category": row[2],
        "batch": row[3], "year": row[4], "city": row[5],
        "district": row[6], "unit": row[7]
    })

items_by_name = {}
for i, it in enumerate(HERITAGE["items"]):
    items_by_name.setdefault(it["name"], []).append(it)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "项目明细"

headers = ["序号", "项目名称", "类别", "批次", "年份", "市", "县/区", "保护单位", "图片"]
ws.append(headers)

hdr_font = Font(bold=True, size=11, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="305496")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col_idx, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col_idx)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = center
    c.border = border

widths = [6, 32, 12, 10, 8, 10, 14, 40, 14]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

THUMB_W = 100
THUMB_H = 75
ROW_HEIGHT = 60

for idx, rd in enumerate(rows_data):
    row_num = idx + 2
    ws.append([rd["seq"], rd["name"], rd["category"], rd["batch"],
               rd["year"], rd["city"], rd["district"], rd["unit"], ""])
    for col in range(1, 10):
        c = ws.cell(row=row_num, column=col)
        c.border = border
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row_num].height = ROW_HEIGHT

    its = items_by_name.get(rd["name"], [])
    if its:
        pid = str(its[0]["id"])
        img_path = os.path.join(IMG_DIR, f"{pid}.jpg")
        if os.path.exists(img_path):
            try:
                pil = PILImage.open(img_path)
                pil.thumbnail((THUMB_W, THUMB_H))
                tmp_path = os.path.join(IMG_DIR, f"_thumb_{pid}.png")
                pil.save(tmp_path, "PNG")
                xi = XLImage(tmp_path)
                xi.width = THUMB_W
                xi.height = THUMB_H
                cell_addr = f"I{row_num}"
                ws.add_image(xi, cell_addr)
            except Exception as e:
                ws.cell(row=row_num, column=9).value = f"[图片错误:{e}]"

ws.freeze_panes = "A2"

ws2 = wb.create_sheet("按市汇总")
city_count = {}
for rd in rows_data:
    city_count[rd["city"]] = city_count.get(rd["city"], 0) + 1
ws2.append(["市", "项目数"])
for c, n in sorted(city_count.items(), key=lambda x: -x[1]):
    ws2.append([c, n])
for col in [1, 2]:
    ws2.cell(row=1, column=col).font = hdr_font
    ws2.cell(row=1, column=col).fill = hdr_fill
ws2.column_dimensions["A"].width = 14
ws2.column_dimensions["B"].width = 10

ws3 = wb.create_sheet("按类别汇总")
cat_count = {}
for rd in rows_data:
    cat_count[rd["category"]] = cat_count.get(rd["category"], 0) + 1
ws3.append(["类别", "项目数"])
for c, n in sorted(cat_count.items(), key=lambda x: -x[1]):
    ws3.append([c, n])
for col in [1, 2]:
    ws3.cell(row=1, column=col).font = hdr_font
    ws3.cell(row=1, column=col).fill = hdr_fill
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 10

wb.save(OUT)
print(f"Saved: {OUT}")
print(f"Rows: {len(rows_data)}, images embedded.")

for f in os.listdir(IMG_DIR):
    if f.startswith("_thumb_"):
        try:
            os.remove(os.path.join(IMG_DIR, f))
        except Exception:
            pass
